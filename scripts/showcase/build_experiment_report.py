#!/usr/bin/env python3
"""Build local, editable experiment deliverables from a passing real demo."""
from __future__ import annotations

import argparse
from datetime import datetime, timezone
import hashlib
from html import escape
from html.parser import HTMLParser
import json
import math
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
from urllib.parse import quote, unquote, urlsplit
import zipfile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE as RT


REPO = Path(__file__).resolve().parents[2]
TITLE = "Qt × Agent 视频创作工作台实验报告"
SOURCE_FILES = [
    "app/ui/MainWindow.cpp", "app/ui/AgentPanel.cpp",
    "app/workflow/AgentWorkbenchBridge.cpp", "app/agent/AgentController.cpp",
    "app/agent/ContextBuilder.cpp", "app/agent/ModelClient.cpp",
    "app/agent/ApprovalManager.h", "app/agent/ToolDispatcher.cpp",
    "app/controllers/EditorController.cpp", "app/controllers/ExportController.cpp",
    "app/services/MediaValidation.cpp", "app/services/SampleCatalog.cpp",
    "tests/integration/AgentWorkbenchDemo.cpp", "templates/video-story/template.json",
    "templates/video-story/packages/video-story/render.js", "config/version-lock.json",
    "docs/EXPERIMENT_REPORT.md", "scripts/showcase/build_experiment_report.py",
]


def require(condition, message):
    if not condition:
        raise ValueError(message)


def read_json(path):
    return json.loads(Path(path).read_text(encoding="utf-8"))


def write_json(path, value):
    Path(path).write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def digest(path):
    h = hashlib.sha256()
    with Path(path).open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def contained_file(base, relative):
    require(isinstance(relative, str) and relative, "缺少相对文件路径")
    path = (base / relative).resolve()
    require(not Path(relative).is_absolute() and path.is_relative_to(base.resolve()),
            f"文件路径越界：{relative}")
    require(path.is_file(), f"文件不存在：{path}")
    return path


def run(command, timeout=300):
    result = subprocess.run(command, capture_output=True, text=True, timeout=timeout)
    require(result.returncode == 0, f"命令失败：{command[0]}\n{result.stderr[-3000:]}")
    return result.stdout


def verify_demo(demo):
    require(demo.get("verdict") == "PASS", "仅接受真实演示 verdict=PASS 的证据")
    for key in ("sourceUnchangedBeforeApproval", "artifactFullDecode", "recordingFullDecode", "ownedStudioStopped"):
        require(demo.get(key) is True, f"演示证据未通过：{key}")
    require(demo.get("pixelsUploaded") is False, "本报告要求演示未上传图像像素")
    require(demo.get("provider") == "current Codex login", "演示需使用当前 Codex 登录")
    require(demo.get("realModelRequests") == 2 and demo.get("explicitDriverApprovals") == 2,
            "演示需实际完成两次模型请求和两次显式批准")
    require(demo.get("recordingFrames", 0) >= 10 and demo.get("recordingElapsedMs", 0) > 0,
            "缺少真实录制帧数或经过时间")
    require(isinstance(demo.get("stages"), list) and len(demo["stages"]) >= 9, "演示阶段记录不完整")
    for key in ("materialPlan", "titlePlan"):
        require(len(demo.get(key, {}).get("operations", [])) == 1, f"{key} 应包含实际单项修改")
    previous = -1
    for stage in demo["stages"]:
        at = stage.get("atMs", -1)
        index = stage.get("frameIndex", -1)
        require(isinstance(at, (int, float)) and previous <= at <= demo["recordingElapsedMs"], "阶段时间无效")
        require(isinstance(index, int) and 0 <= index < demo["recordingFrames"], "阶段帧号无效")
        require(isinstance(stage.get("title"), str), "缺少阶段标题")
        previous = at


def copy_video(source, target, title, role, provenance, ffprobe, ffmpeg, expected_hash=None, verified_decode=None):
    source = Path(source).resolve()
    require(source.is_file(), f"缺少视频：{source}")
    shutil.copy2(source, target)
    sha = digest(target)
    require(sha == digest(source), f"视频复制后哈希不一致：{source}")
    if expected_hash:
        require(sha == expected_hash, f"视频与来源清单哈希不一致：{source}")
    info = json.loads(run([ffprobe, "-v", "error", "-show_format", "-show_streams", "-of", "json", str(target)]))
    videos = [s for s in info.get("streams", []) if s.get("codec_type") == "video"]
    require(videos, f"缺少视频流：{source}")
    stream = videos[0]
    duration = float(info.get("format", {}).get("duration", 0))
    require(math.isfinite(duration) and duration > 0, f"视频时长无效：{source}")
    require(stream.get("width", 0) > 0 and stream.get("height", 0) > 0, f"视频尺寸无效：{source}")
    # A producer's boolean alone does not bind its decode to these bytes.
    # Reuse only evidence accompanied by an expected hash checked above.
    reused_decode = verified_decode if expected_hash else None
    if not reused_decode:
        run([ffmpeg, "-v", "error", "-xerror", "-i", str(target), "-f", "null", "-"], timeout=600)
    return {"title": title, "role": role, "path": f"videos/{target.name}",
            "sourcePath": str(source), "provenance": provenance, "sha256": sha,
            "sizeBytes": target.stat().st_size, "durationSeconds": duration,
            "width": stream["width"], "height": stream["height"], "codec": stream.get("codec_name"),
            "frameRate": stream.get("avg_frame_rate"),
            "audioStreams": sum(s.get("codec_type") == "audio" for s in info["streams"]),
            "validation": {"ffprobeExitCode": 0, "fullDecodeExitCode": 0,
                           "fullDecodeEvidence": reused_decode or "delivery copy decoded by report generator",
                           "copiedBytesMatchSource": True}, "ffprobe": info}


def format_time(seconds):
    return f"{int(seconds) // 60:02d}:{int(seconds) % 60:02d}"


def report_sections(demo, media, prep, lock, generated):
    film = media[1]
    duplicate_count = len(prep.get("existingLocalSamples", []))
    unique_count = len({x["sha256"] for x in prep.get("existingLocalSamples", [])})
    return [
        ("实验目的", [
            "实现可运行的 Qt 6 Widgets 视频创作工作台，贯通工程管理、素材选择、自然语言提案、显式审阅批准、预览同步和本地视频导出。重点检验 Qt 操作与 Agent 的实际协作，以及修改权限和导出成功语义。",
            "本次实验由 Qt Test 脚本驱动生产窗口控件与信号，共用应用正式使用的 AgentWorkbenchBridge；模型为当前 Codex CLI 登录会话，编译、预览及渲染使用真实本地 Hypit。窗口录像来自本应用 MainWindow，保留实际帧间隔和模型等待时间；不是手工操作录像。",
        ]),
        ("系统架构", [
            "Qt UI（MainWindow / AgentPanel） → AgentWorkbenchBridge → AgentController → ContextBuilder / ModelClient → Codex JSON 提案 → PlanReview / ApprovalManager → ToolDispatcher / EditorController → Hypit Studio；导出由 ExportController 顺序执行 plan、build、get、ffprobe 与全片解码。",
            "原创部分包括 Qt 窗口、工程与素材管理、模板、属性编辑、Agent 上下文与提案校验、版本/批准约束、异步任务状态及交付验收。复用部分包括 Hypit 视频语言、编译与渲染、Studio 预览服务、Qt 框架、Codex CLI 和 FFmpeg；官方示例视频单独保留来源。",
            f"版本依据：仓库 config/version-lock.json 固定 Qt {lock['qt']['version']}、Hypit {lock['hypit']['version']}、C++ {lock['toolchain']['cxxStandard']}，平台记录为 {lock['platform']['os']} {lock['platform']['osVersion']} / {lock['platform']['arch']}。这些是项目固定版本记录；本次媒体验证工具版本另见 source-manifest.json。报告生成时间：{generated}。",
        ]),
        ("Qt 与 Agent 联动实现", [
            "在素材列表中选择一个可见且兼容的已登记素材后，“交给 Agent 调整”将真实素材绑定值写入目标输入，并限定到当前组件。多选、不兼容或任务忙碌等状态禁止交接。交接只准备目标；生成方案与批准执行分别由按钮触发。",
            "AgentWorkbenchBridge 连接 generateRequested、scopeChanged、selectionChanged 和 approveRequested 等信号。AgentController 把当前工程、真实 entityId/fieldId、可写属性、版本指纹及登记素材交给 ContextBuilder。ModelClient 获取符合 JSON Schema 的提案；模型不直接执行本地编辑命令。",
            "方案进入审阅后展示拟修改内容。ApprovalManager 管理批准，执行链再次校验方案、允许范围和版本，再由 ToolDispatcher / EditorController 写入 Hypit。实验断言批准前源码未变，并在执行后等待预览指纹与新源码一致。",
            "本轮上下文只提供素材名称、类型、尺寸与绑定等结构信息，imagePixelsProvided=false；没有上传图片像素，也不将提案描述解释为模型看过视频画面。",
            "ExportController 冻结当前输入，先验证本地运行计划，再核对 Build 状态和精确 Output；只有媒体参数验证及全片解码完成后才进入 complete。付费渲染 Provider 请求不通过导出计划门禁。",
        ]),
        ("操作步骤", [
            "1. 准备固定版本 Hypit、Qt、FFmpeg 与当前已登录 Codex CLI；生成并验证本地官方素材目录，应用启动读取本地 catalog.json。",
            "2. 查看示例库，从本地 Hypit 对话样例创建 video-story 工程，等待 Studio 会话与预览确认，并在 Qt 内播放。",
            "3. 导入官方排行榜视频，选中兼容组件和素材，点击“交给 Agent 调整”，检查目标中出现实际素材路径。",
            "4. 点击生成方案，等待真实 Codex 响应。审阅素材替换提案，确认源码尚未改变，再显式批准并等待新版本预览。",
            "5. 输入“只把当前视频组件标题改为‘Qt 与 Agent，让创作连起来’，其它文字、视频素材、时长和颜色都保持不变。”再次生成、审阅并批准。",
            "6. 导出成片，等待计划、构建、输出获取和媒体解码完成，检查任务与工作历史。运行报告脚本，把验证后的素材、录像、成片与证据复制到交付目录。",
        ]),
        ("测试结果", [
            f"最终演示结论：{demo['verdict']}，包含 {demo['realModelRequests']} 次真实模型请求和 {demo['explicitDriverApprovals']} 次脚本显式批准。这是最终演示的计数，不代表开发及重录期间的请求总量。批准前源码不变、成片全片解码、窗口录像全片解码及自有 Studio 关闭均由演示断言通过。Build ID：{demo['buildId']}。",
            f"交付脚本对复制后的全部 {len(media)} 个视频重新计算 SHA-256 并运行 ffprobe；对完整操作录像、工程成片和本地对话样例重新执行全片解码，四个官方视频则通过预先记录的 SHA-256 绑定其已有解码证据。实际成片为 {film['durationSeconds']:.3f} 秒、{film['width']} × {film['height']}，包含 {film['audioStreams']} 条音频流。video-story 模板设计为 8 秒静音视频卡。",
            f"来源清单中的上游本地视频共 {duplicate_count} 个文件，对应 {unique_count} 个不同 SHA-256。新增官方素材按哈希验证为 {len(media) - 3} 个不同视频；短于 8 秒的访谈片段仅供示例播放，不能直接满足当前视频模板时长。",
            "本报告呈现这一条真实演示流程的断言与媒体校验结果；不据此声称所有用户输入、模型回答或全部平台均已验证。完整自动化测试集的结果应查阅仓库独立测试证据。",
            "失败与恢复边界：模型失败、澄清或过期方案不会自动批准；Studio 409 冲突与 422 回滚按既有恢复流程处理；导出各阶段失败不能显示成功。本次 PASS 记录未包含故障注入，不宣称这些异常路径在本次录像中被逐一演示。",
            "录制修正：首轮窗口采集保留了 Retina 图像的设备像素比，导致窗口画面只占编码画布的一部分。采集器在缩放后归一化设备像素比，并重新运行完整真实演示。交付视频取修正后的录制，首轮模型请求不计入上面的最终演示次数。",
        ]),
        ("演示视频", [
            f"完整窗口录制：{media[0]['durationSeconds']:.3f} 秒，{demo['recordingFrames']} 个原始采集帧，实际采集经过 {demo['recordingElapsedMs'] / 1000:.3f} 秒。每个阶段时间来自演示 JSON；编码成片与采集时间可能因最后一帧保持和编码取整略有差异。",
            "本节提供完整操作录制、导出成片及官方/本地示例。HTML 视频直接读取随报告提供的 videos 文件夹，无需联网；Word 中的视频链接指向同一相对路径，建议保留整个交付目录一起移动。",
            "官方示例由 Hypit 提供，用于展示与本地实验，不计入本工作台原创视频成果。source-manifest.json 记录官方 URL、压缩包成员、原始路径、哈希与本次验证参数；工程输出的源素材归属仍属于其原来源。",
        ]),
        ("总结", [
            "实验把 Qt 的素材操作与自然语言 Agent 请求连接到同一条可审阅执行链，实际完成素材替换、标题修改、预览确认和本地成片导出。源码指纹与显式批准共同约束执行时机，导出验证区分了命令接受、构建完成、输出获取和视频可解码。",
            "当前限制：模型依赖现有 Codex 登录及响应可用性；本轮为单一工程的两次受限编辑；不提供视频视觉理解证据；视频模板固定取前 8 秒且静音。完整录像由程序驱动真实 UI，不用于证明人工可用性评测。后续可增加更多实际任务、异常恢复实验以及长视频与音频模板。",
        ]),
    ]


def add_link(paragraph, text, target):
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), paragraph.part.relate_to(target, RT.HYPERLINK, is_external=True))
    run_element = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "166A83")
    props.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(underline)
    run_element.append(props)
    value = OxmlElement("w:t")
    value.text = text
    run_element.append(value)
    hyperlink.append(run_element)
    paragraph._p.append(hyperlink)


def build_docx(target, sections, demo, media, images):
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = Inches(.75)
    for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2"):
        style = doc.styles[name]
        style.font.name = "Arial"
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "等线")
        style.font.size = Pt(11 if name == "Normal" else 15 if name.startswith("Heading") else 25)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.25
    doc.styles["Normal"].paragraph_format.space_after = Pt(8)
    doc.core_properties.title = TITLE
    doc.core_properties.subject = "真实 Qt / Codex / Hypit 集成实验"
    doc.core_properties.author = ""
    doc.add_heading(TITLE, 0)
    doc.add_paragraph("真实流程 · 显式审阅 · 本地成片", "Subtitle")
    doc.add_paragraph("课程：____________    姓名：____________    学号：____________")
    doc.add_paragraph("以上信息由提交人填写；实验结果依据随附机器证据生成。")
    for number, (title, paragraphs) in enumerate(sections, 1):
        doc.add_heading(f"{number}. {title}", 1)
        for paragraph in paragraphs:
            doc.add_paragraph(paragraph)
        if title == "测试结果":
            table = doc.add_table(rows=1, cols=3)
            table.style = "Light Shading Accent 1"
            for cell, label in zip(table.rows[0].cells, ("验证项", "结果", "证据")):
                cell.text = label
            for row in result_rows(demo):
                for cell, value in zip(table.add_row().cells, row):
                    cell.text = str(value)
        if title == "演示视频":
            for item in media:
                paragraph = doc.add_paragraph()
                add_link(paragraph, item["title"], quote(item["path"]))
                paragraph.add_run(f"  |  {item['durationSeconds']:.3f} 秒 · {item['width']}×{item['height']} · {item['role']}")
            doc.add_heading("阶段时间表", 2)
            for stage in demo["stages"]:
                doc.add_paragraph(f"{format_time(stage['atMs'] / 1000)}  {stage['title']}")
            for image in images:
                doc.add_picture(str(target.parent / image["path"]), width=Inches(6.0))
                doc.add_paragraph(f"真实窗口帧 {image['frameIndex']}：{image['title']}")
    doc.add_heading("附录：来源与复现", 1)
    doc.add_paragraph("源代码说明见仓库 docs/EXPERIMENT_REPORT.md。官方素材溯源、版本记录、逐文件校验值及实际操作提案见随附 JSON。")
    for label, path in (("离线报告入口", "index.html"), ("交付来源与校验清单", "source-manifest.json"), ("真实演示证据", "evidence/agent-workbench-demo.json")):
        add_link(doc.add_paragraph(), label, path)
    footer = sec.footer.paragraphs[0]
    footer.text = "Qt × Agent 实验报告 · 依据真实演示证据生成"
    footer.runs[0].font.size = Pt(9)
    footer.runs[0].font.color.rgb = RGBColor.from_string("667085")
    doc.save(target)
    with zipfile.ZipFile(target) as archive:
        require(archive.testzip() is None, "Word ZIP 完整性验证失败")
    Document(target)


def result_rows(demo):
    return [
        ("真实集成演示", demo["verdict"], "agent-workbench-demo.json"),
        ("Codex 请求 / 显式批准", f"{demo['realModelRequests']} / {demo['explicitDriverApprovals']}", "materialPlan、titlePlan"),
        ("批准前源码不变", "通过", "sourceUnchangedBeforeApproval"),
        ("像素上传", "未上传", "pixelsUploaded=false"),
        ("实际编译版本", demo["previewFingerprint"], "previewFingerprint"),
        ("成片与录像全片解码", "通过", "演示断言 + 交付副本重新解码"),
        ("自有 Studio 已关闭", "通过", "ownedStudioStopped"),
    ]


def build_html(target, sections, demo, media, images):
    nav = "".join(f'<a href="#section-{i}">{escape(title)}</a>' for i, (title, _) in enumerate(sections, 1))
    body = []
    for i, (title, paragraphs) in enumerate(sections, 1):
        content = "".join(f"<p>{escape(text)}</p>" for text in paragraphs)
        if title == "系统架构":
            content += '<div class="flow">Qt 操作 → Agent 上下文 → Codex 提案 → 审阅批准 → Hypit 修改 → 预览确认 → 导出校验</div>'
        if title == "测试结果":
            content += '<div class="table-wrap"><table><thead><tr><th>验证项</th><th>结果</th><th>证据</th></tr></thead><tbody>'
            content += "".join("<tr>" + "".join(f"<td>{escape(str(value))}</td>" for value in row) + "</tr>" for row in result_rows(demo))
            content += "</tbody></table></div>"
        if title == "演示视频":
            content += '<div class="videos">'
            for index, item in enumerate(media):
                source = item["provenance"]
                url = source.get("sourceUrl")
                attribution = f'<a href="{escape(url, quote=True)}">Hypit 官方素材来源</a>' if url else escape(source["description"])
                content += f'''<article class="video-card {'wide' if index == 0 else ''}"><h3>{escape(item['title'])}</h3>
<video controls preload="metadata" playsinline src="{quote(item['path'])}"></video>
<p>{item['durationSeconds']:.3f} 秒 · {item['width']} × {item['height']} · {escape(item['role'])}</p>
<p class="source">{attribution}</p><a href="{quote(item['path'])}" download>下载视频</a></article>'''
            content += '</div><h3>完整录像阶段时间表</h3><ol class="timeline">'
            content += "".join(f"<li><time>{format_time(s['atMs'] / 1000)}</time> {escape(s['title'])}</li>" for s in demo["stages"])
            content += "</ol>"
            if images:
                content += '<h3>真实窗口帧</h3><div class="frames">' + "".join(
                    f'<figure><a href="{quote(im["path"])}"><img loading="lazy" src="{quote(im["path"])}" alt="{escape(im["title"], quote=True)}"></a><figcaption>采集帧 {im["frameIndex"]} · {escape(im["title"])}</figcaption></figure>' for im in images) + '</div>'
        body.append(f'<section id="section-{i}"><span class="section-no">0{i}</span><h2>{escape(title)}</h2>{content}</section>')
    target.write_text('''<!doctype html><html lang="zh-CN"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1"><title>''' + TITLE + '''</title><style>
:root{color-scheme:light;--ink:#192c3a;--muted:#5b6d79;--teal:#126b70;--line:#dbe4e8}*{box-sizing:border-box}html{scroll-behavior:smooth}body{margin:0;font-family:-apple-system,BlinkMacSystemFont,"PingFang SC","Microsoft YaHei",sans-serif;background:#f1f5f7;color:var(--ink);line-height:1.85}header{background:#122938;color:#fff;padding:66px max(24px,calc((100vw - 1080px)/2)) 48px}.eyebrow{color:#83dbcd;letter-spacing:.18em;font-size:13px}h1{font-size:clamp(30px,4vw,48px);line-height:1.25;margin:18px 0}header p{max-width:760px;color:#c6d5dc}.badges{display:flex;gap:12px;flex-wrap:wrap}.badges span{background:#244451;border:1px solid #476470;border-radius:24px;padding:5px 16px}nav{display:flex;flex-wrap:wrap;gap:10px;padding:20px 0}nav a,.download{color:#c7ffef}main{max-width:1120px;margin:24px auto;padding:0 20px}section{position:relative;background:#fff;border:1px solid var(--line);border-radius:16px;padding:32px 38px;margin-bottom:22px;scroll-margin-top:20px}h2{font-size:26px;margin:0 0 18px}.section-no{float:right;color:#8aabb7;font-size:30px}p{margin:12px 0}a{color:var(--teal);overflow-wrap:anywhere}code,td{overflow-wrap:anywhere}.flow{padding:22px;background:#eef7f4;border-left:4px solid #258d81;border-radius:8px;font-weight:600;margin-top:24px}.table-wrap{overflow:auto}table{width:100%;border-collapse:collapse;font-size:14px;margin-top:24px}th,td{border-bottom:1px solid var(--line);text-align:left;padding:12px;max-width:390px}th{background:#f1f6f8}.videos,.frames{display:grid;grid-template-columns:1fr 1fr;gap:22px;margin:24px 0}.video-card{border:1px solid var(--line);border-radius:12px;padding:18px;background:#f7fafb}.video-card.wide{grid-column:1/-1}.video-card h3{margin:0 0 12px}video{display:block;width:100%;max-height:520px;background:#0e1620;border-radius:8px}.source,figcaption{font-size:13px;color:var(--muted)}.timeline{padding-left:24px}.timeline li{padding:6px}.timeline time{font-variant-numeric:tabular-nums;color:var(--teal);font-weight:600;margin-right:12px}figure{margin:0}img{width:100%;border-radius:8px;border:1px solid var(--line)}footer{padding:24px 0 48px;color:var(--muted);font-size:14px}@media(max-width:650px){header{padding:38px 22px}section{padding:24px 20px}.videos,.frames{grid-template-columns:1fr}}@media print{body{background:white}header{color:#192c3a;background:white;padding:20px}header p{color:#192c3a}nav,video{display:none}section{break-inside:avoid}.videos{display:block}}
</style></head><body><header><div class="eyebrow">FRAMELAB · 实验交付</div><h1>''' + TITLE + '''</h1><p>从真实素材到可审阅方案，再到实际成片。以生产窗口、当前 Codex 登录与本地 Hypit 执行结果为依据。</p><div class="badges"><span>真实集成 PASS</span><span>2 次模型请求</span><span>显式批准后执行</span></div><nav>''' + nav + '''</nav><a class="download" href="experiment-report.docx" download>下载可编辑 Word 报告 ↗</a></header><main>''' + "".join(body) + '''<footer><p>交付文件：<a href="experiment-report.docx">Word 报告</a> · <a href="source-manifest.json">来源与校验清单</a> · <a href="evidence/agent-workbench-demo.json">真实演示 JSON</a> · <a href="evidence/catalog.json">官方素材目录</a> · <a href="evidence/preparation-manifest.json">素材准备证据</a> · <a href="evidence/EXPERIMENT_REPORT.md">复现说明</a></p><p>离线打开本页即可播放随附视频。课程、姓名和学号请在 Word 中自行填写。</p></footer></main></body></html>''', encoding="utf-8")


class Links(HTMLParser):
    def __init__(self):
        super().__init__()
        self.targets = []
        self.ids = set()

    def handle_starttag(self, tag, attrs):
        for key, value in attrs:
            if key in ("href", "src") and value:
                self.targets.append(value)
            elif key == "id":
                self.ids.add(value)


def validate_links(folder):
    parser = Links()
    parser.feed((folder / "index.html").read_text(encoding="utf-8"))
    targets = list(parser.targets)
    doc = Document(folder / "experiment-report.docx")
    targets += [r.target_ref for r in doc.part.rels.values() if r.reltype == RT.HYPERLINK]
    for target in targets:
        parsed = urlsplit(target)
        if parsed.scheme in ("https", "http"):
            continue
        require(not parsed.scheme and not parsed.netloc, f"未知链接类型：{target}")
        if parsed.path:
            contained_file(folder, unquote(parsed.path))
        if parsed.fragment:
            require(unquote(parsed.fragment) in parser.ids, f"失效章节链接：{target}")
    return len(targets)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--delivery", type=Path, default=REPO / ".workbench/deliverables-m13")
    parser.add_argument("--demo", type=Path, default=REPO / "docs/evidence/m13/agent-workbench-demo.json")
    parser.add_argument("--catalog", type=Path, default=REPO / ".workbench/showcase/catalog.json")
    parser.add_argument("--preparation", type=Path, default=REPO / ".workbench/showcase/preparation-manifest.json")
    parser.add_argument("--ffprobe", default=shutil.which("ffprobe"))
    parser.add_argument("--ffmpeg", default=shutil.which("ffmpeg"))
    args = parser.parse_args()
    require(args.ffprobe and args.ffmpeg, "需要本地 ffprobe 与 ffmpeg")
    demo, catalog, prep = read_json(args.demo), read_json(args.catalog), read_json(args.preparation)
    verify_demo(demo)
    require(prep.get("status") == "verified" and not prep.get("failures"), "官方素材准备尚未全部通过")
    entries = catalog.get("samples", [])
    require(len(entries) == 4 and len({x["sha256"] for x in entries}) == 4, "需要四个不同的官方视频")
    require(entries == prep.get("samples"), "catalog 与素材准备清单不一致")
    for entry in entries:
        require(entry.get("validation", {}).get("fullDecodeExitCode") == 0, "官方素材缺少完整验证")
        require(urlsplit(entry.get("sourceUrl", "")).scheme == "https", "官方素材缺少 HTTPS 来源")
    local_samples = prep.get("existingLocalSamples", [])
    local_entry = next((e for e in local_samples if e["path"] == "output/chat-demo.mp4"), None)
    require(local_entry is not None, "缺少本地对话样例来源")
    lock = read_json(REPO / "config/version-lock.json")
    delivery = args.delivery.resolve()
    delivery.mkdir(parents=True, exist_ok=True)
    generated = datetime.now(timezone.utc).isoformat()
    tool_versions = {"ffprobe": run([args.ffprobe, "-version"]).splitlines()[0],
                     "ffmpeg": run([args.ffmpeg, "-version"]).splitlines()[0]}
    with tempfile.TemporaryDirectory(prefix=".report-build-", dir=delivery) as temporary:
        stage = Path(temporary)
        for child in ("videos", "evidence", "images"):
            (stage / child).mkdir()
        media = []

        def video(source, filename, title, role, provenance, expected_hash=None, verified_decode=None):
            print(f"验证交付视频：{title}", flush=True)
            item = copy_video(source, stage / "videos" / filename, title, role, provenance,
                              args.ffprobe, args.ffmpeg, expected_hash, verified_decode)
            media.append(item)

        video(demo["recording"], "qt-agent-walkthrough.mp4", "Qt × Agent 完整操作录制", "程序驱动真实窗口；保留等待时间",
              {"description": demo["driver"], "capture": demo["capture"], "evidence": "evidence/agent-workbench-demo.json"},
              verified_decode="evidence/agent-workbench-demo.json: recordingFullDecode=true")
        video(demo["artifact"], "qt-agent-film.mp4", "本次工程导出成片", "8 秒静音视频卡",
              {"description": "本工作台原创 video-story 模板；素材复用 Hypit 官方排行榜示例", "buildId": demo["buildId"], "sourceProject": demo["project"]},
              verified_decode="evidence/agent-workbench-demo.json: artifactFullDecode=true")
        require(abs(media[1]["durationSeconds"] - 8) < .1, "成片与本次 8 秒模板描述不符")
        require(abs(media[0]["durationSeconds"] - demo["recordingElapsedMs"] / 1000) < 2,
                "录制时长与实际采集经过不一致")
        for index, entry in enumerate(entries, 1):
            source = contained_file(args.catalog.resolve().parent, entry["path"])
            video(source, f"official-{index:02d}.mp4", entry["name"],
                  "官方示例；仅预览（不足 8 秒）" if entry["durationSeconds"] < 8 else "官方示例素材",
                  {key: entry[key] for key in ("sourceUrl", "archiveMember") if key in entry}, entry["sha256"],
                  "evidence/preparation-manifest.json: matching sample SHA-256 and fullDecodeExitCode=0")
        local = contained_file(Path(prep["distribution"]).resolve(), local_entry["path"])
        video(local, "local-chat.mp4", "本地 Hypit 对话样例", "复用本地上游样例；原有重复文件仅计一份",
              {"description": "固定版本 Hypit 本地 output/chat-demo.mp4", "upstreamCommit": prep["upstreamCommit"], "upstreamPath": local_entry["path"]}, local_entry["sha256"])
        images = []
        for number in (4, 5, 8):
            record = demo["stages"][number]
            frame = Path(demo["recording"]).parent / "frames" / f"frame-{record['frameIndex']:06d}.jpg"
            if frame.is_file():
                name = f"stage-{number + 1:02d}.jpg"
                shutil.copy2(frame, stage / "images" / name)
                images.append({"title": record["title"], "path": f"images/{name}", "frameIndex": record["frameIndex"],
                               "stageAtMs": record["atMs"], "sourcePath": str(frame), "sha256": digest(stage / "images" / name)})
        for source, name in ((args.demo, "agent-workbench-demo.json"), (args.catalog, "catalog.json"),
                             (args.preparation, "preparation-manifest.json"), (REPO / "config/version-lock.json", "version-lock.json"),
                             (REPO / "docs/EXPERIMENT_REPORT.md", "EXPERIMENT_REPORT.md")):
            shutil.copy2(source, stage / "evidence" / name)
        sections = report_sections(demo, media, prep, lock, generated)
        build_docx(stage / "experiment-report.docx", sections, demo, media, images)
        build_html(stage / "index.html", sections, demo, media, images)
        manifest = {"format": "qvw.experiment-delivery@1", "generatedAt": generated, "verdict": "PASS",
                    "scope": "真实演示证据与本次交付媒体验证；不代表全测试集结果", "tools": tool_versions,
                    "media": media, "images": images, "officialSources": prep,
                    "sourceFiles": [{"path": p, "sha256": digest(REPO / p)} for p in SOURCE_FILES],
                    "files": [{"path": str(p.relative_to(stage)), "sha256": digest(p), "sizeBytes": p.stat().st_size}
                              for p in sorted(stage.rglob("*")) if p.is_file()]}
        write_json(stage / "source-manifest.json", manifest)
        manifest["validatedLinks"] = validate_links(stage)
        write_json(stage / "source-manifest.json", manifest)
        for child in stage.iterdir():
            target = delivery / child.name
            require(not target.is_symlink(), f"交付目标不得为符号链接：{target}")
            if child.is_dir():
                target.mkdir(exist_ok=True)
                for source in child.iterdir():
                    source.replace(target / source.name)
            else:
                child.replace(target)
    validate_links(delivery)
    print(json.dumps({"verdict": "PASS", "delivery": str(delivery), "videos": len(media),
                      "images": len(images), "validatedLinks": manifest["validatedLinks"]}, ensure_ascii=False))


if __name__ == "__main__":
    try:
        main()
    except (ValueError, OSError, KeyError, TypeError, subprocess.TimeoutExpired) as error:
        print(f"报告生成失败：{error}", file=sys.stderr)
        sys.exit(1)
